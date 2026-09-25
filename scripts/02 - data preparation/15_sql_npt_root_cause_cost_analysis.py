# ================================================================
# STAGE 2G.7 — SQL NPT ROOT CAUSE & COST ANALYSIS
# Project: Guyana Offshore Drilling Analytics
# ================================================================
#
# Purpose:
#   Analyze Non-Productive Time (NPT) using the SQL data model.
#
# Main analyses:
#   1. NPT category analysis
#   2. NPT subcategory analysis
#   3. Severity analysis
#   4. Well-level NPT ranking
#   5. Rig-level NPT ranking
#   6. Well × Rig analysis
#   7. NPT duration analysis
#   8. NPT cost analysis
#   9. Pareto analysis
#  10. Overall KPI summary
#
# Output:
#   outputs/sql_stage_2G.7_NPT_Root_Cause_Cost_Analysis.xlsx
#
# ================================================================

import sqlite3
from pathlib import Path
import pandas as pd
import numpy as np


# ================================================================
# 1. PROJECT PATHS
# ================================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

DATA_DIR = PROJECT_ROOT / "data"
DB_DIR = DATA_DIR / "sql"
OUTPUT_DIR = PROJECT_ROOT / "outputs"

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


# ================================================================
# 2. SQLITE DATABASE
# ================================================================

DB_PATH = PROJECT_ROOT / "database" / "guyana_drilling.db"

if not DB_PATH.exists():

    print("=" * 72)
    print("ERROR — SQLITE DATABASE NOT FOUND")
    print("=" * 72)

    print(f"\nExpected database:")
    print(f"  {DB_PATH}")

    raise FileNotFoundError(
        f"SQLite database not found: {DB_PATH}"
    )


# ================================================================
# 3. OUTPUT FILE
# ================================================================

OUTPUT_FILE = (
    OUTPUT_DIR /
    "sql_stage_2G.7_NPT_Root_Cause_Cost_Analysis.xlsx"
)


# ================================================================
# 4. HEADER
# ================================================================

print("=" * 72)
print("STAGE 2G.7 — SQL NPT ROOT CAUSE & COST ANALYSIS")
print("=" * 72)

print(f"\nProject root : {PROJECT_ROOT}")
print(f"Database     : {DB_PATH}")
print(f"Output       : {OUTPUT_FILE}")


# ================================================================
# 5. CONNECT TO DATABASE
# ================================================================

conn = sqlite3.connect(DB_PATH)


# ================================================================
# 6. DATABASE INSPECTION
# ================================================================

tables = pd.read_sql_query(
    """
    SELECT name
    FROM sqlite_master
    WHERE type = 'table'
    ORDER BY name
    """,
    conn
)

print("\nAvailable SQL tables:")

for table in tables["name"]:
    print(f"  - {table}")


required_tables = ["Fact_NPT"]

missing_tables = [
    table
    for table in required_tables
    if table not in tables["name"].tolist()
]

if missing_tables:

    conn.close()

    print("\nERROR — Missing required table(s):")

    for table in missing_tables:
        print(f"  - {table}")

    raise RuntimeError(
        "Fact_NPT is required for Stage 2G.7."
    )


# ================================================================
# 7. INSPECT FACT_NPT STRUCTURE
# ================================================================

print("\nInspecting Fact_NPT structure...")

fact_columns = pd.read_sql_query(
    "PRAGMA table_info(Fact_NPT)",
    conn
)

print("\nFact_NPT columns:")

for column in fact_columns["name"]:
    print(f"  - {column}")


available_columns = fact_columns["name"].tolist()


# ================================================================
# 8. IDENTIFY COLUMN NAMES
# ================================================================

def find_column(possible_names, available):

    for name in possible_names:

        if name in available:
            return name

    return None


date_col = find_column(
    ["Date", "NPT_Date", "Event_Date"],
    available_columns
)

well_col = find_column(
    ["Well_ID", "WellID"],
    available_columns
)

rig_col = find_column(
    ["Rig_ID", "RigID"],
    available_columns
)

category_col = find_column(
    ["NPT_Category", "Category"],
    available_columns
)

subcategory_col = find_column(
    ["NPT_Subcategory", "Subcategory"],
    available_columns
)

duration_col = find_column(
    ["Duration_hr", "NPT_Duration_hr", "Duration_Hours"],
    available_columns
)

cost_col = find_column(
    ["Cost_USD", "NPT_Cost_USD", "Cost"],
    available_columns
)

severity_col = find_column(
    ["Severity", "NPT_Severity"],
    available_columns
)


print("\nDetected columns:")

print(f"  Date          : {date_col}")
print(f"  Well          : {well_col}")
print(f"  Rig           : {rig_col}")
print(f"  Category      : {category_col}")
print(f"  Subcategory   : {subcategory_col}")
print(f"  Duration      : {duration_col}")
print(f"  Cost          : {cost_col}")
print(f"  Severity      : {severity_col}")


# ================================================================
# 9. REQUIRED NPT COLUMNS
# ================================================================

required_fact_columns = {
    "well": well_col,
    "rig": rig_col,
    "category": category_col,
    "subcategory": subcategory_col,
    "duration": duration_col,
    "cost": cost_col,
    "severity": severity_col,
}

missing_columns = [
    key
    for key, value in required_fact_columns.items()
    if value is None
]

if missing_columns:

    conn.close()

    print("\nERROR — Required Fact_NPT columns were not found:")

    for column in missing_columns:
        print(f"  - {column}")

    raise RuntimeError(
        "Fact_NPT structure does not contain the expected analytical fields."
    )


# ================================================================
# 10. LOAD FACT_NPT
# ================================================================

print("\nLoading Fact_NPT...")

fact_npt = pd.read_sql_query(
    "SELECT * FROM Fact_NPT",
    conn
)

print(f"\nFact_NPT rows loaded: {len(fact_npt):,}")


if fact_npt.empty:

    conn.close()

    raise RuntimeError(
        "Fact_NPT is empty. Generate Fact_NPT before running Stage 2G.7."
    )


# ================================================================
# 11. NORMALIZE DATA
# ================================================================

fact_npt[duration_col] = pd.to_numeric(
    fact_npt[duration_col],
    errors="coerce"
)

fact_npt[cost_col] = pd.to_numeric(
    fact_npt[cost_col],
    errors="coerce"
)

fact_npt[duration_col] = fact_npt[duration_col].fillna(0)

fact_npt[cost_col] = fact_npt[cost_col].fillna(0)


# ================================================================
# 12. BASIC DATA QUALITY
# ================================================================

data_quality = pd.DataFrame({
    "Metric": [
        "Total NPT Events",
        "Null Well IDs",
        "Null Rig IDs",
        "Null Categories",
        "Null Subcategories",
        "Null Severity",
        "Null Duration",
        "Null Cost",
        "Negative Duration",
        "Negative Cost",
    ],

    "Value": [
        len(fact_npt),
        fact_npt[well_col].isna().sum(),
        fact_npt[rig_col].isna().sum(),
        fact_npt[category_col].isna().sum(),
        fact_npt[subcategory_col].isna().sum(),
        fact_npt[severity_col].isna().sum(),
        fact_npt[duration_col].isna().sum(),
        fact_npt[cost_col].isna().sum(),
        (fact_npt[duration_col] < 0).sum(),
        (fact_npt[cost_col] < 0).sum(),
    ]
})


# ================================================================
# 13. OVERALL KPIs
# ================================================================

total_events = len(fact_npt)

total_npt_hours = fact_npt[duration_col].sum()

total_npt_cost = fact_npt[cost_col].sum()

average_duration = (
    fact_npt[duration_col].mean()
    if total_events > 0
    else 0
)

average_cost = (
    fact_npt[cost_col].mean()
    if total_events > 0
    else 0
)

max_duration = fact_npt[duration_col].max()

max_cost = fact_npt[cost_col].max()

unique_wells = fact_npt[well_col].nunique()

unique_rigs = fact_npt[rig_col].nunique()


overall_kpis = pd.DataFrame({
    "KPI": [
        "Total NPT Events",
        "Total NPT Hours",
        "Total NPT Cost USD",
        "Average NPT Duration hr",
        "Average NPT Cost USD",
        "Maximum NPT Duration hr",
        "Maximum NPT Cost USD",
        "Wells Affected",
        "Rigs Affected",
    ],

    "Value": [
        total_events,
        total_npt_hours,
        total_npt_cost,
        average_duration,
        average_cost,
        max_duration,
        max_cost,
        unique_wells,
        unique_rigs,
    ]
})


# ================================================================
# 14. CATEGORY ANALYSIS
# ================================================================

category_analysis = (
    fact_npt
    .groupby(category_col, dropna=False)
    .agg(
        NPT_Events=(category_col, "size"),
        NPT_Hours=(duration_col, "sum"),
        NPT_Cost_USD=(cost_col, "sum"),
        Avg_Duration_hr=(duration_col, "mean"),
        Avg_Cost_USD=(cost_col, "mean"),
    )
    .reset_index()
)


category_analysis["Pct_Total_Hours"] = (
    category_analysis["NPT_Hours"]
    / total_npt_hours
    * 100
    if total_npt_hours > 0
    else 0
)


category_analysis["Pct_Total_Cost"] = (
    category_analysis["NPT_Cost_USD"]
    / total_npt_cost
    * 100
    if total_npt_cost > 0
    else 0
)


category_analysis = category_analysis.sort_values(
    "NPT_Cost_USD",
    ascending=False
).reset_index(drop=True)


category_analysis["Cumulative_Cost_Pct"] = (
    category_analysis["Pct_Total_Cost"].cumsum()
)


category_analysis["Cost_Rank"] = (
    category_analysis["NPT_Cost_USD"]
    .rank(method="dense", ascending=False)
    .astype(int)
)


# ================================================================
# 15. SUBCATEGORY ANALYSIS
# ================================================================

subcategory_analysis = (
    fact_npt
    .groupby(
        [category_col, subcategory_col],
        dropna=False
    )
    .agg(
        NPT_Events=(subcategory_col, "size"),
        NPT_Hours=(duration_col, "sum"),
        NPT_Cost_USD=(cost_col, "sum"),
        Avg_Duration_hr=(duration_col, "mean"),
        Avg_Cost_USD=(cost_col, "mean"),
    )
    .reset_index()
)


subcategory_analysis["Pct_Total_Hours"] = (
    subcategory_analysis["NPT_Hours"]
    / total_npt_hours
    * 100
    if total_npt_hours > 0
    else 0
)


subcategory_analysis["Pct_Total_Cost"] = (
    subcategory_analysis["NPT_Cost_USD"]
    / total_npt_cost
    * 100
    if total_npt_cost > 0
    else 0
)


subcategory_analysis = subcategory_analysis.sort_values(
    "NPT_Cost_USD",
    ascending=False
).reset_index(drop=True)


subcategory_analysis["Cumulative_Cost_Pct"] = (
    subcategory_analysis["Pct_Total_Cost"].cumsum()
)


subcategory_analysis["Cost_Rank"] = (
    subcategory_analysis["NPT_Cost_USD"]
    .rank(method="dense", ascending=False)
    .astype(int)
)


# ================================================================
# 16. SEVERITY ANALYSIS
# ================================================================

severity_analysis = (
    fact_npt
    .groupby(severity_col, dropna=False)
    .agg(
        NPT_Events=(severity_col, "size"),
        NPT_Hours=(duration_col, "sum"),
        NPT_Cost_USD=(cost_col, "sum"),
        Avg_Duration_hr=(duration_col, "mean"),
        Avg_Cost_USD=(cost_col, "mean"),
    )
    .reset_index()
)


severity_analysis["Pct_Total_Events"] = (
    severity_analysis["NPT_Events"]
    / total_events
    * 100
    if total_events > 0
    else 0
)


severity_analysis["Pct_Total_Hours"] = (
    severity_analysis["NPT_Hours"]
    / total_npt_hours
    * 100
    if total_npt_hours > 0
    else 0
)


severity_analysis["Pct_Total_Cost"] = (
    severity_analysis["NPT_Cost_USD"]
    / total_npt_cost
    * 100
    if total_npt_cost > 0
    else 0
)


severity_order = {
    "Low": 1,
    "Medium": 2,
    "High": 3
}


severity_analysis["_sort"] = (
    severity_analysis[severity_col]
    .map(severity_order)
    .fillna(99)
)


severity_analysis = (
    severity_analysis
    .sort_values("_sort")
    .drop(columns="_sort")
    .reset_index(drop=True)
)


# ================================================================
# 17. WELL-LEVEL ANALYSIS
# ================================================================

well_analysis = (
    fact_npt
    .groupby(well_col, dropna=False)
    .agg(
        NPT_Events=(well_col, "size"),
        NPT_Hours=(duration_col, "sum"),
        NPT_Cost_USD=(cost_col, "sum"),
        Avg_Duration_hr=(duration_col, "mean"),
        Avg_Cost_USD=(cost_col, "mean"),
        Max_Duration_hr=(duration_col, "max"),
        Max_Cost_USD=(cost_col, "max"),
        NPT_Categories=(category_col, "nunique"),
        NPT_Subcategories=(subcategory_col, "nunique"),
    )
    .reset_index()
)


well_analysis["Pct_Total_Cost"] = (
    well_analysis["NPT_Cost_USD"]
    / total_npt_cost
    * 100
    if total_npt_cost > 0
    else 0
)


well_analysis["Cost_Rank"] = (
    well_analysis["NPT_Cost_USD"]
    .rank(method="dense", ascending=False)
    .astype(int)
)


well_analysis = well_analysis.sort_values(
    "NPT_Cost_USD",
    ascending=False
).reset_index(drop=True)


# ================================================================
# 18. RIG-LEVEL ANALYSIS
# ================================================================

rig_analysis = (
    fact_npt
    .groupby(rig_col, dropna=False)
    .agg(
        NPT_Events=(rig_col, "size"),
        NPT_Hours=(duration_col, "sum"),
        NPT_Cost_USD=(cost_col, "sum"),
        Avg_Duration_hr=(duration_col, "mean"),
        Avg_Cost_USD=(cost_col, "mean"),
        Max_Duration_hr=(duration_col, "max"),
        Max_Cost_USD=(cost_col, "max"),
        Wells_Affected=(well_col, "nunique"),
        NPT_Categories=(category_col, "nunique"),
    )
    .reset_index()
)


rig_analysis["Pct_Total_Cost"] = (
    rig_analysis["NPT_Cost_USD"]
    / total_npt_cost
    * 100
    if total_npt_cost > 0
    else 0
)


rig_analysis["Cost_Rank"] = (
    rig_analysis["NPT_Cost_USD"]
    .rank(method="dense", ascending=False)
    .astype(int)
)


rig_analysis = rig_analysis.sort_values(
    "NPT_Cost_USD",
    ascending=False
).reset_index(drop=True)


# ================================================================
# 19. WELL × RIG ANALYSIS
# ================================================================

well_rig_analysis = (
    fact_npt
    .groupby(
        [well_col, rig_col],
        dropna=False
    )
    .agg(
        NPT_Events=(well_col, "size"),
        NPT_Hours=(duration_col, "sum"),
        NPT_Cost_USD=(cost_col, "sum"),
        Avg_Duration_hr=(duration_col, "mean"),
        Avg_Cost_USD=(cost_col, "mean"),
    )
    .reset_index()
)


well_rig_analysis["Pct_Total_Cost"] = (
    well_rig_analysis["NPT_Cost_USD"]
    / total_npt_cost
    * 100
    if total_npt_cost > 0
    else 0
)


well_rig_analysis = well_rig_analysis.sort_values(
    "NPT_Cost_USD",
    ascending=False
).reset_index(drop=True)


# ================================================================
# 20. CATEGORY × SEVERITY ANALYSIS
# ================================================================

category_severity = (
    fact_npt
    .groupby(
        [category_col, severity_col],
        dropna=False
    )
    .agg(
        NPT_Events=(category_col, "size"),
        NPT_Hours=(duration_col, "sum"),
        NPT_Cost_USD=(cost_col, "sum"),
        Avg_Duration_hr=(duration_col, "mean"),
        Avg_Cost_USD=(cost_col, "mean"),
    )
    .reset_index()
)


category_severity["Pct_Total_Cost"] = (
    category_severity["NPT_Cost_USD"]
    / total_npt_cost
    * 100
    if total_npt_cost > 0
    else 0
)


category_severity = category_severity.sort_values(
    "NPT_Cost_USD",
    ascending=False
).reset_index(drop=True)


# ================================================================
# 21. SUBCATEGORY × SEVERITY
# ================================================================

subcategory_severity = (
    fact_npt
    .groupby(
        [
            category_col,
            subcategory_col,
            severity_col
        ],
        dropna=False
    )
    .agg(
        NPT_Events=(subcategory_col, "size"),
        NPT_Hours=(duration_col, "sum"),
        NPT_Cost_USD=(cost_col, "sum"),
        Avg_Duration_hr=(duration_col, "mean"),
        Avg_Cost_USD=(cost_col, "mean"),
    )
    .reset_index()
)


subcategory_severity["Pct_Total_Cost"] = (
    subcategory_severity["NPT_Cost_USD"]
    / total_npt_cost
    * 100
    if total_npt_cost > 0
    else 0
)


subcategory_severity = subcategory_severity.sort_values(
    "NPT_Cost_USD",
    ascending=False
).reset_index(drop=True)


# ================================================================
# 22. DURATION BUCKET ANALYSIS
# ================================================================

def duration_bucket(hours):

    if hours <= 4:
        return "0–4 hr"

    elif hours <= 8:
        return "4–8 hr"

    elif hours <= 12:
        return "8–12 hr"

    elif hours <= 24:
        return "12–24 hr"

    else:
        return ">24 hr"


fact_npt["Duration_Bucket"] = (
    fact_npt[duration_col]
    .apply(duration_bucket)
)


duration_analysis = (
    fact_npt
    .groupby("Duration_Bucket", dropna=False)
    .agg(
        NPT_Events=("Duration_Bucket", "size"),
        NPT_Hours=(duration_col, "sum"),
        NPT_Cost_USD=(cost_col, "sum"),
        Avg_Duration_hr=(duration_col, "mean"),
        Avg_Cost_USD=(cost_col, "mean"),
    )
    .reset_index()
)


duration_order = {
    "0–4 hr": 1,
    "4–8 hr": 2,
    "8–12 hr": 3,
    "12–24 hr": 4,
    ">24 hr": 5
}


duration_analysis["_sort"] = (
    duration_analysis["Duration_Bucket"]
    .map(duration_order)
)


duration_analysis = (
    duration_analysis
    .sort_values("_sort")
    .drop(columns="_sort")
    .reset_index(drop=True)
)


duration_analysis["Pct_Total_Cost"] = (
    duration_analysis["NPT_Cost_USD"]
    / total_npt_cost
    * 100
    if total_npt_cost > 0
    else 0
)


# ================================================================
# 23. HIGH-SEVERITY NPT
# ================================================================

high_severity = fact_npt[
    fact_npt[severity_col]
    .astype(str)
    .str.upper()
    .eq("HIGH")
].copy()


high_severity_analysis = (
    high_severity
    .groupby(category_col, dropna=False)
    .agg(
        High_Severity_Events=(category_col, "size"),
        High_Severity_Hours=(duration_col, "sum"),
        High_Severity_Cost_USD=(cost_col, "sum"),
        Avg_Duration_hr=(duration_col, "mean"),
        Avg_Cost_USD=(cost_col, "mean"),
    )
    .reset_index()
)


high_severity_analysis["Pct_High_Severity_Cost"] = (
    high_severity_analysis["High_Severity_Cost_USD"]
    / high_severity_analysis["High_Severity_Cost_USD"].sum()
    * 100
    if len(high_severity_analysis) > 0
    and high_severity_analysis["High_Severity_Cost_USD"].sum() > 0
    else 0
)


high_severity_analysis = high_severity_analysis.sort_values(
    "High_Severity_Cost_USD",
    ascending=False
).reset_index(drop=True)


# ================================================================
# 24. TOP NPT EVENTS
# ================================================================

top_events = (
    fact_npt
    .sort_values(
        [cost_col, duration_col],
        ascending=[False, False]
    )
    .head(25)
    .copy()
)


top_events.insert(
    0,
    "Cost_Rank",
    range(1, len(top_events) + 1)
)


# ================================================================
# 25. MONTHLY NPT TREND
# ================================================================

if date_col is not None:

    fact_npt["_Analysis_Date"] = pd.to_datetime(
        fact_npt[date_col],
        errors="coerce"
    )

    fact_npt["_Year_Month"] = (
        fact_npt["_Analysis_Date"]
        .dt.to_period("M")
        .astype(str)
    )

    monthly_analysis = (
        fact_npt
        .groupby("_Year_Month", dropna=False)
        .agg(
            NPT_Events=(category_col, "size"),
            NPT_Hours=(duration_col, "sum"),
            NPT_Cost_USD=(cost_col, "sum"),
            Avg_Duration_hr=(duration_col, "mean"),
            Avg_Cost_USD=(cost_col, "mean"),
        )
        .reset_index()
        .rename(columns={"_Year_Month": "Year_Month"})
    )

    monthly_analysis["Pct_Total_Cost"] = (
        monthly_analysis["NPT_Cost_USD"]
        / total_npt_cost
        * 100
        if total_npt_cost > 0
        else 0
    )

else:

    monthly_analysis = pd.DataFrame({
        "Year_Month": [],
        "NPT_Events": [],
        "NPT_Hours": [],
        "NPT_Cost_USD": [],
        "Avg_Duration_hr": [],
        "Avg_Cost_USD": [],
        "Pct_Total_Cost": [],
    })


# ================================================================
# 26. TOP ROOT CAUSES
# ================================================================

top_root_causes = (
    subcategory_analysis
    .head(10)
    .copy()
)


# ================================================================
# 27. PARETO 80/20 FLAG
# ================================================================

category_analysis["Pareto_80_Flag"] = np.where(
    category_analysis["Cumulative_Cost_Pct"] <= 80,
    "Within 80%",
    "Outside 80%"
)


subcategory_analysis["Pareto_80_Flag"] = np.where(
    subcategory_analysis["Cumulative_Cost_Pct"] <= 80,
    "Within 80%",
    "Outside 80%"
)


# ================================================================
# 28. EXECUTIVE SUMMARY
# ================================================================

if not category_analysis.empty:

    top_category = category_analysis.iloc[0][category_col]
    top_category_cost = category_analysis.iloc[0]["NPT_Cost_USD"]

else:

    top_category = "N/A"
    top_category_cost = 0


if not subcategory_analysis.empty:

    top_subcategory = (
        subcategory_analysis.iloc[0][subcategory_col]
    )

    top_subcategory_cost = (
        subcategory_analysis.iloc[0]["NPT_Cost_USD"]
    )

else:

    top_subcategory = "N/A"
    top_subcategory_cost = 0


if not well_analysis.empty:

    worst_well = well_analysis.iloc[0][well_col]
    worst_well_cost = well_analysis.iloc[0]["NPT_Cost_USD"]

else:

    worst_well = "N/A"
    worst_well_cost = 0


if not rig_analysis.empty:

    worst_rig = rig_analysis.iloc[0][rig_col]
    worst_rig_cost = rig_analysis.iloc[0]["NPT_Cost_USD"]

else:

    worst_rig = "N/A"
    worst_rig_cost = 0


executive_summary = pd.DataFrame({

    "Metric": [
        "Total NPT Events",
        "Total NPT Hours",
        "Total NPT Cost USD",
        "Average NPT Duration hr",
        "Average NPT Cost USD",
        "Top NPT Category",
        "Top NPT Category Cost USD",
        "Top NPT Subcategory",
        "Top NPT Subcategory Cost USD",
        "Highest NPT Cost Well",
        "Highest NPT Cost Well USD",
        "Highest NPT Cost Rig",
        "Highest NPT Cost Rig USD",
        "Wells Affected",
        "Rigs Affected",
    ],

    "Value": [
        total_events,
        total_npt_hours,
        total_npt_cost,
        average_duration,
        average_cost,
        top_category,
        top_category_cost,
        top_subcategory,
        top_subcategory_cost,
        worst_well,
        worst_well_cost,
        worst_rig,
        worst_rig_cost,
        unique_wells,
        unique_rigs,
    ]
})


# ================================================================
# 29. PRINT RESULTS
# ================================================================

print("\n" + "=" * 72)
print("EXECUTIVE NPT SUMMARY")
print("=" * 72)

print(f"\nTotal NPT Events       : {total_events:,}")
print(f"Total NPT Hours        : {total_npt_hours:,.2f}")
print(f"Total NPT Cost         : ${total_npt_cost:,.2f}")
print(f"Average NPT Duration   : {average_duration:,.2f} hr")
print(f"Average NPT Cost       : ${average_cost:,.2f}")

print(f"\nTop Category           : {top_category}")
print(f"Top Category Cost      : ${top_category_cost:,.2f}")

print(f"\nTop Subcategory        : {top_subcategory}")
print(f"Top Subcategory Cost   : ${top_subcategory_cost:,.2f}")

print(f"\nHighest Cost Well      : {worst_well}")
print(f"Highest Well NPT Cost  : ${worst_well_cost:,.2f}")

print(f"\nHighest Cost Rig       : {worst_rig}")
print(f"Highest Rig NPT Cost   : ${worst_rig_cost:,.2f}")


# ================================================================
# 30. CATEGORY TABLE
# ================================================================

print("\n" + "=" * 72)
print("NPT CATEGORY ANALYSIS")
print("=" * 72)

print(
    category_analysis[
        [
            category_col,
            "NPT_Events",
            "NPT_Hours",
            "NPT_Cost_USD",
            "Pct_Total_Cost",
        ]
    ].to_string(index=False)
)


# ================================================================
# 31. TOP ROOT CAUSES
# ================================================================

print("\n" + "=" * 72)
print("TOP 10 NPT ROOT CAUSES BY COST")
print("=" * 72)

print(
    top_root_causes[
        [
            category_col,
            subcategory_col,
            "NPT_Events",
            "NPT_Hours",
            "NPT_Cost_USD",
            "Pct_Total_Cost",
        ]
    ].to_string(index=False)
)


# ================================================================
# 32. TOP WELLS
# ================================================================

print("\n" + "=" * 72)
print("TOP 10 WELLS BY NPT COST")
print("=" * 72)

print(
    well_analysis[
        [
            well_col,
            "NPT_Events",
            "NPT_Hours",
            "NPT_Cost_USD",
            "Pct_Total_Cost",
        ]
    ].head(10).to_string(index=False)
)


# ================================================================
# 33. TOP RIGS
# ================================================================

print("\n" + "=" * 72)
print("RIG NPT ANALYSIS")
print("=" * 72)

print(
    rig_analysis[
        [
            rig_col,
            "NPT_Events",
            "NPT_Hours",
            "NPT_Cost_USD",
            "Wells_Affected",
            "Pct_Total_Cost",
        ]
    ].to_string(index=False)
)


# ================================================================
# 34. EXPORT TO EXCEL
# ================================================================

print("\n" + "=" * 72)
print("EXPORTING RESULTS")
print("=" * 72)


with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    executive_summary.to_excel(
        writer,
        sheet_name="Executive_Summary",
        index=False
    )

    overall_kpis.to_excel(
        writer,
        sheet_name="Overall_KPIs",
        index=False
    )

    category_analysis.to_excel(
        writer,
        sheet_name="NPT_Category",
        index=False
    )

    subcategory_analysis.to_excel(
        writer,
        sheet_name="NPT_Subcategory",
        index=False
    )

    severity_analysis.to_excel(
        writer,
        sheet_name="Severity",
        index=False
    )

    well_analysis.to_excel(
        writer,
        sheet_name="Well_Analysis",
        index=False
    )

    rig_analysis.to_excel(
        writer,
        sheet_name="Rig_Analysis",
        index=False
    )

    well_rig_analysis.to_excel(
        writer,
        sheet_name="Well_x_Rig",
        index=False
    )

    category_severity.to_excel(
        writer,
        sheet_name="Category_x_Severity",
        index=False
    )

    subcategory_severity.to_excel(
        writer,
        sheet_name="Subcategory_x_Severity",
        index=False
    )

    duration_analysis.to_excel(
        writer,
        sheet_name="Duration_Buckets",
        index=False
    )

    high_severity_analysis.to_excel(
        writer,
        sheet_name="High_Severity",
        index=False
    )

    top_events.to_excel(
        writer,
        sheet_name="Top_25_Events",
        index=False
    )

    monthly_analysis.to_excel(
        writer,
        sheet_name="Monthly_Trend",
        index=False
    )

    data_quality.to_excel(
        writer,
        sheet_name="Data_Quality",
        index=False
    )


# ================================================================
# 35. FORMAT EXCEL
# ================================================================

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


wb = load_workbook(OUTPUT_FILE)


header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

header_font = Font(
    bold=True,
    color="FFFFFF"
)


for ws in wb.worksheets:

    # Header formatting
    for cell in ws[1]:

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center"
        )

    # Freeze header
    ws.freeze_panes = "A2"

    # Autofit columns
    for column_cells in ws.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            try:

                value_length = len(
                    str(cell.value)
                )

                max_length = max(
                    max_length,
                    value_length
                )

            except Exception:

                pass

        ws.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 10),
            40
        )


wb.save(OUTPUT_FILE)


# ================================================================
# 36. CLOSE DATABASE
# ================================================================

conn.close()


# ================================================================
# 37. FINAL STATUS
# ================================================================

print("\n" + "=" * 72)
print("STAGE 2G.7 COMPLETED SUCCESSFULLY")
print("=" * 72)

print(f"\nOutput file:")
print(f"  {OUTPUT_FILE}")

print("\nGenerated analysis sheets:")

sheets = [
    "Executive_Summary",
    "Overall_KPIs",
    "NPT_Category",
    "NPT_Subcategory",
    "Severity",
    "Well_Analysis",
    "Rig_Analysis",
    "Well_x_Rig",
    "Category_x_Severity",
    "Subcategory_x_Severity",
    "Duration_Buckets",
    "High_Severity",
    "Top_25_Events",
    "Monthly_Trend",
    "Data_Quality",
]

for sheet in sheets:
    print(f"  ✓ {sheet}")

print("\nStage 2G.7 is ready for validation.")
print("=" * 72)